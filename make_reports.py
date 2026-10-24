"""Create Word reports for practical works 5 and 6."""

from docx import Document
from docx.shared import Pt
from pathlib import Path


OUT = Path(__file__).parent / "reports"


def report(filename, title, sections):
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)
    document.add_heading(title, 0)
    document.add_paragraph("ФИО: ______________________________________________")
    document.add_paragraph("Группа: ____________________    Преподаватель: ____________________")
    for heading, text in sections:
        document.add_heading(heading, level=1)
        for paragraph in text.split("\n\n"):
            document.add_paragraph(paragraph)
    document.save(OUT / filename)


report("Отчет_ПР5.docx", "Отчёт по практической работе №5\nDjango, HTTP и основы веб-разработки", [
    ("Цель работы", "Освоить создание Django-проекта и приложения, маршрутизацию URL, view-функции и формирование HTML-страниц.") ,
    ("Выполненная работа", "Созданы manage.py, пакет visitor_project и приложение webapp. В settings.py зарегистрировано приложение, настроены русская локализация, шаблоны и static-файлы. В корневом urls.py подключены маршруты webapp через include()."),
    ("Предметная область", "Страницы системы отображают посетителей, сотрудников, пропуска и посещения. Списки и карточки доступны по именованным URL, а динамические параметры ограничены конвертером <int:...>. Для неизвестного идентификатора возвращается 404."),
    ("Источник данных", "Веб-часть продолжает использовать JSON-файлы data/*.json и классы ПР3 через storage.py. Django ORM, миграции, формы и SQLite на ПР5 не используются согласно методичке."),
    ("HTML, CSS и Bootstrap", "Добавлены главная страница, списки и детали, единое оформление CSS и Bootstrap CDN. Все view используют render(request, template, context)."),
    ("Проверка", "Выполнены команды: python manage.py check; python manage.py test webapp; python -m pytest -v; python -m flake8 . --exclude=.venv. Все проверки проходят."),
    ("Вывод", "Создан рабочий Django-интерфейс учёта посетителей с HTTP-маршрутизацией и отображением данных предметной области."),
])

report("Отчет_ПР6.docx", "Отчёт по практической работе №6\nШаблоны Django", [
    ("Цель работы", "Освоить шаблоны Django: наследование, блоки, include, контекст, циклы, условия, фильтры и static-ресурсы."),
    ("Базовый и дочерние шаблоны", "templates/base.html содержит общую HTML-структуру, навигацию и подключения CSS/JavaScript. Дочерние страницы используют extends и переопределяют блоки title и content."),
    ("Переиспользуемые фрагменты", "Навигация вынесена в templates/includes/navigation.html. Карточка посетителя и статус посещения подключаются через include, поэтому разметка не дублируется."),
    ("Работа с данными", "View передают контекст в render. В шаблонах используются обращение к атрибутам объектов, for/empty, if, фильтры length, default, date и capfirst. Именованные URL namespace webapp формируют ссылки через тег url."),
    ("Static и ошибки", "Подключены Bootstrap, собственный CSS, JavaScript и SVG-иллюстрация через load static/static. Добавлен шаблон 404.html. Некорректные и отсутствующие идентификаторы обрабатываются корректно."),
    ("Ограничения текущего этапа", "JSON остаётся источником данных. ORM, база данных и Django Forms не используются: их изучение относится к следующим практикам."),
    ("Проверка", "Выполнены команды: python manage.py check; python manage.py test webapp; python -m pytest -v; python -m flake8 . --exclude=.venv. В браузере проверены главная, списки, детали и 404."),
    ("Вывод", "Система получила единый адаптивный интерфейс на шаблонах Django с наследованием, компонентами и статическими ресурсами."),
])
